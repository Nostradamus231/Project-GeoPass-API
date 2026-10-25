import io
import json
import urllib.error
import urllib.request
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_URL = "http://127.0.0.1:8000"


def request_http(url: str, data: dict = None, method: str = "GET", raw_bytes: bytes = None):
    headers = {"Content-Type": "application/json; charset=utf-8"}
    if raw_bytes is not None:
        body = raw_bytes
    elif data is not None:
        body = json.dumps(data).encode("utf-8")
    else:
        body = None
    req = urllib.request.Request(url, data=body, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            status = resp.status
            content = resp.read()
            cd_hdr = resp.headers.get("content-disposition", "")
            try:
                parsed = json.loads(content.decode("utf-8"))
            except Exception:
                parsed = content
            return status, parsed, cd_hdr
    except urllib.error.HTTPError as e:
        status = e.code
        content = e.read()
        cd_hdr = e.headers.get("content-disposition", "")
        try:
            parsed = json.loads(content.decode("utf-8"))
        except Exception:
            parsed = content.decode("utf-8", errors="replace")
        return status, parsed, cd_hdr


def get_base_payload() -> dict:
    return {
        "company_name": "ТОО Тестовый Карьер",
        "deposit_name": "Северное",
        "license_number": "ЛИЦ-001",
        "mineral_type": "Песок строительный",
        "region": "Костанайская область",
        "coordinate_system": "local_metric",
        "contour": [
            {"x": 0.0, "y": 0.0},
            {"x": 100.0, "y": 0.0},
            {"x": 100.0, "y": 100.0},
            {"x": 0.0, "y": 100.0},
        ],
        "boreholes": [
            {"borehole_id": "S-1", "x": 10.0, "y": 10.0, "thickness_m": 5.0, "weight": 1.0},
            {"borehole_id": "S-2", "x": 90.0, "y": 10.0, "thickness_m": 5.0, "weight": 1.0},
            {"borehole_id": "S-3", "x": 90.0, "y": 90.0, "thickness_m": 5.0, "weight": 1.0},
            {"borehole_id": "S-4", "x": 10.0, "y": 90.0, "thickness_m": 5.0, "weight": 1.0},
        ],
        "quality": {
            "modulus_fineness": 2.1,
            "clay_content_pct": 3.5,
            "moisture_pct": 4.0,
            "bulk_density_t_m3": 1.5,
        },
        "density_t_m3": 1.6,
        "declared_area_ha": 1.0,
    }


def run_checks() -> dict[str, tuple[bool, str]]:
    results = {}

    st1, body1, _ = request_http(f"{BASE_URL}/generate-passport", data=get_base_payload(), method="POST")
    dl_url = body1.get("download_url", "")
    st_dl, doc_bytes, cd1 = request_http(f"{BASE_URL}{dl_url}", method="GET")
    doc = Document(io.BytesIO(doc_bytes))

    leftovers = []
    for p in doc.paragraphs:
        if "{{" in p.text or "{%" in p.text:
            leftovers.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if "{{" in c.text or "{%" in c.text:
                    leftovers.append(c.text)
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            if "{{" in p.text or "{%" in p.text:
                leftovers.append(p.text)
        for p in sec.footer.paragraphs:
            if "{{" in p.text or "{%" in p.text:
                leftovers.append(p.text)

    check1_passed = st1 == 200 and st_dl == 200 and len(leftovers) == 0
    results["1. Генерация и скачивание"] = (check1_passed, f"Status: {st1}, leftovers: {len(leftovers)}")

    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                full_text.append(c.text)
    joined_text = " ".join(full_text)
    has_raw_numbers = "80000.0" in joined_text or "10000.0" in joined_text or "50000.0" in joined_text
    has_formatted_80k = "80\u00a0000" in joined_text or "80 000" in joined_text
    has_formatted_50k = "50\u00a0000" in joined_text or "50 000" in joined_text
    has_formatted_area = "1,00" in joined_text
    has_formatted_thick = "5,00" in joined_text

    check2_passed = not has_raw_numbers and has_formatted_80k and has_formatted_50k and has_formatted_area and has_formatted_thick
    results["2. Числа в документе"] = (check2_passed, f"Raw floats: {has_raw_numbers}, 80k: {has_formatted_80k}, 1.00: {has_formatted_area}")

    t_geo = doc.tables[2]
    t_res = doc.tables[5]
    geo_indicators = [r.cells[0].text.strip() for r in t_geo.rows[1:]]
    res_indicators = [r.cells[0].text.strip() for r in t_res.rows[1:]]

    check3_passed = len(geo_indicators) == len(set(geo_indicators)) and len(res_indicators) == len(set(res_indicators))
    results["3. Дублирование показателей"] = (check3_passed, f"Geo unique: {len(geo_indicators) == len(set(geo_indicators))}, Res unique: {len(res_indicators) == len(set(res_indicators))}")

    t_bh = doc.tables[3]
    bh_rows_count = len(t_bh.rows)
    data_rows = t_bh.rows[1:]
    alignments_ok = True
    for r in data_rows:
        for idx in [1, 2, 3, 4]:
            p = r.cells[idx].paragraphs[0]
            if p.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                alignments_ok = False

    check4_passed = bh_rows_count == 5 and alignments_ok
    results["4. Таблица скважин"] = (check4_passed, f"Rows: {bh_rows_count}, right-aligned: {alignments_ok}")

    sec = doc.sections[0]
    different_first_page = sec.different_first_page_header_footer
    footer_text = "".join(p.text for p in sec.footer.paragraphs)
    footer_has_page = "Страница" in footer_text and "из" in footer_text

    check5_passed = different_first_page and footer_has_page
    results["5. Колонтитулы"] = (check5_passed, f"Different first page: {different_first_page}, Footer text: {footer_text}")

    p6 = get_base_payload()
    p6["boreholes"] = [
        {"borehole_id": "S-IN", "x": 50.0, "y": 50.0, "thickness_m": 5.0, "weight": 1.0},
        {"borehole_id": "S-OUT", "x": 500.0, "y": 500.0, "thickness_m": 5.0, "weight": 1.0},
    ]
    st6, body6, _ = request_http(f"{BASE_URL}/generate-passport", data=p6, method="POST")
    dl_url6 = body6.get("download_url", "")
    st_dl6, bytes6, _ = request_http(f"{BASE_URL}{dl_url6}", method="GET")
    doc6 = Document(io.BytesIO(bytes6))
    outside_list = body6.get("summary", {}).get("boreholes_outside_contour", [])

    check6_passed = st6 == 200 and st_dl6 == 200 and "S-OUT" in outside_list and len(doc6.paragraphs) > 0
    results["6. Скважина вне контура"] = (check6_passed, f"Status: {st6}, outside: {outside_list}")

    long_title = "Северный_Перспективный_Участок_Месторождения_Песков_Строительных_И_Гравия_Номер_Один"
    p7 = get_base_payload()
    p7["deposit_name"] = long_title[:80]
    st7, body7, _ = request_http(f"{BASE_URL}/generate-passport", data=p7, method="POST")
    dl_url7 = body7.get("download_url", "")
    st_dl7, bytes7, cd7 = request_http(f"{BASE_URL}{dl_url7}", method="GET")
    doc7 = Document(io.BytesIO(bytes7))

    check7_passed = st7 == 200 and st_dl7 == 200 and len(doc7.paragraphs) > 0 and len(doc7.tables) == 6
    results["7. Длинное название и кириллица"] = (check7_passed, f"Status: {st7}, cd: {cd7[:40]}..., doc tables: {len(doc7.tables)}")

    p8_no_company = get_base_payload()
    del p8_no_company["company_name"]
    st8_missing, body8_missing, _ = request_http(f"{BASE_URL}/generate-passport", data=p8_no_company, method="POST")
    fields_missing = [e.get("field") for e in body8_missing.get("errors", [])]

    p8_neg = get_base_payload()
    p8_neg["boreholes"][0]["thickness_m"] = -5.0
    st8_neg, _, _ = request_http(f"{BASE_URL}/generate-passport", data=p8_neg, method="POST")

    p8_bf = get_base_payload()
    p8_bf["contour"] = [
        {"x": 0.0, "y": 0.0},
        {"x": 100.0, "y": 100.0},
        {"x": 100.0, "y": 0.0},
        {"x": 0.0, "y": 100.0},
    ]
    st8_bf, body8_bf, _ = request_http(f"{BASE_URL}/generate-passport", data=p8_bf, method="POST")
    bf_msg = json.dumps(body8_bf, ensure_ascii=False)

    st8_bad_json, body8_bad_json, _ = request_http(f"{BASE_URL}/generate-passport", raw_bytes=b"{\"broken\": ", method="POST")
    bad_json_msg = json.dumps(body8_bad_json, ensure_ascii=False)

    st8_trav, _, _ = request_http(f"{BASE_URL}/download/..%2Fmain.py", method="GET")
    st8_none, _, _ = request_http(f"{BASE_URL}/download/00000000-0000-0000-0000-000000000000", method="GET")

    check8_passed = (
        st8_missing == 422
        and "company_name" in fields_missing
        and st8_neg == 422
        and st8_bf == 422
        and "самопересекается" in bf_msg
        and st8_bad_json == 422
        and "Тело запроса не является корректным JSON" in bad_json_msg
        and st8_trav == 404
        and st8_none == 404
    )
    results["8. Регрессия API"] = (check8_passed, f"Missing 422: {st8_missing == 422}, Neg 422: {st8_neg == 422}, Butterfly 422: {st8_bf == 422}, Bad JSON 422: {st8_bad_json == 422}, Traversal 404: {st8_trav == 404}")

    return results


if __name__ == "__main__":
    test_results = run_checks()
    all_ok = True
    for name, (ok, detail) in test_results.items():
        status_str = "PASS" if ok else "FAIL"
        if not ok:
            all_ok = False
        print(f"[{status_str}] {name}: {detail}")
    print("SUMMARY:", "ALL_PASSED" if all_ok else "SOME_FAILED")
