from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


def _set_cell_background(cell, hex_color: str):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_padding(cell, top: int = 100, bottom: int = 100, left: int = 140, right: int = 140):
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(tcMar)


def _set_table_borders(table, style: str = "grid"):
    if style == "horizontal":
        b_xml = (
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E6ED"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
    elif style == "box":
        b_xml = (
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:left w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:right w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:insideH w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
    else:
        b_xml = (
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="B0C4DE"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="B0C4DE"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E6ED"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E6ED"/>'
            f'</w:tblBorders>'
        )
    table._tbl.tblPr.append(parse_xml(b_xml))


def _set_col_widths(table, widths):
    for row in table.rows:
        for idx, w in enumerate(widths):
            cell = row.cells[idx]
            cell.width = w
            dxa = int(w.pt * 20)
            tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{dxa}" w:type="dxa"/>')
            cell._tc.get_or_add_tcPr().append(tcW)


def _format_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def generate_template(output_path: Path) -> None:
    doc = Document()

    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(30)
    sec.right_margin = Mm(20)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(20)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1_style = doc.styles["Heading 1"]
    h1_style.font.name = "Times New Roman"
    h1_style.font.size = Pt(13)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(6)
    h1_style.paragraph_format.line_spacing = 1.15
    h1_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(10)
    title_style.paragraph_format.space_after = Pt(4)

    p_hdr = sec.header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_hdr.add_run("Паспорт «Н» • Объект: {{ deposit_name }}")
    r_hdr.font.name = "Times New Roman"
    r_hdr.font.size = Pt(9)
    r_hdr.font.italic = True
    r_hdr.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    p_ftr = sec.footer.paragraphs[0]
    p_ftr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_f1 = p_ftr.add_run("Страница ")
    r_f1.font.name = "Times New Roman"
    r_f1.font.size = Pt(9)
    r_f1.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    fld1 = OxmlElement("w:fldSimple")
    fld1.set(qn("w:instr"), "PAGE")
    p_ftr._p.append(fld1)
    r_f2 = p_ftr.add_run(" из ")
    r_f2.font.name = "Times New Roman"
    r_f2.font.size = Pt(9)
    r_f2.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    fld2 = OxmlElement("w:fldSimple")
    fld2.set(qn("w:instr"), "NUMPAGES")
    p_ftr._p.append(fld2)

    p_stamp = doc.add_paragraph()
    p_stamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_stamp.paragraph_format.space_after = Pt(12)
    r_st = p_stamp.add_run("Гриф: {{ fmt.confidentiality_stamp }} | Экз. № {{ fmt.copy_number }}")
    r_st.font.size = Pt(10)
    r_st.font.italic = True
    r_st.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_title = doc.add_paragraph(style="Title")
    p_title.add_run("Паспорт «Н»")

    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_after = Pt(16)
    r_sub = p_subtitle.add_run("общераспространённые полезные ископаемые, техногенные минеральные образования из вскрышных и/или вмещающих пород")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    t_reg = doc.add_table(rows=4, cols=2)
    t_reg.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_reg, style="horizontal")
    reg_widths = [Mm(95), Mm(65)]
    reg_data = [
        ("Номер паспорта территориального подразделения", "{{ fmt.territorial_passport_number }}"),
        ("Номер паспорта уполномоченного органа", "{{ fmt.authority_passport_number }}"),
        ("Наименование объекта (месторождения)", "{{ deposit_name }}"),
        ("Полезное ископаемое", "{{ mineral_type }}"),
    ]
    for idx, (lbl, val) in enumerate(reg_data):
        row = t_reg.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        p0 = row.cells[0].paragraphs[0]
        p0.add_run(lbl).font.bold = True
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(val)
    _set_col_widths(t_reg, reg_widths)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(4)
    r_sp = p_sp.add_run("Приёмка паспорта:")
    r_sp.font.bold = True
    r_sp.font.size = Pt(11)

    t_recv = doc.add_table(rows=3, cols=5)
    t_recv.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_recv, style="grid")
    recv_widths = [Mm(45), Mm(35), Mm(35), Mm(25), Mm(20)]
    _format_header_row(t_recv.rows[0])
    recv_headers = ["Организация", "ФИО", "Должность", "Подпись", "Дата"]
    for i, h in enumerate(recv_headers):
        cell = t_recv.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(h).font.bold = True

    t_recv_rows = [
        "Территориальное подразделение",
        "Уполномоченный орган",
    ]
    for r_idx, org in enumerate(t_recv_rows, start=1):
        row = t_recv.rows[r_idx]
        for c_idx in range(5):
            _set_cell_padding(row.cells[c_idx], top=80, bottom=80, left=80, right=80)
        p0 = row.cells[0].paragraphs[0]
        p0.add_run(org)
    _set_col_widths(t_recv, recv_widths)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(12)
    p_sig.paragraph_format.space_after = Pt(2)
    r_comp = p_sig.add_run("Составитель: ")
    r_comp.font.bold = True
    p_sig.add_run("{{ fmt.compiler_name }}        Подпись: ___________   Дата: {{ fmt.passport_date }}    М.П.\n")
    r_chk = p_sig.add_run("Проверяющий: ")
    r_chk.font.bold = True
    p_sig.add_run("{{ fmt.checker_name }}        Подпись: ___________   Дата: ___________")

    doc.add_page_break()

    doc.add_heading("Раздел 1. Географические координаты объекта", level=1)
    p_c_note = doc.add_paragraph()
    p_c_note.add_run("{{ fmt.corner_points_note }}")

    t_geo = doc.add_table(rows=4, cols=7)
    t_geo.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_geo, style="grid")
    geo_widths = [Mm(16), Mm(24), Mm(24), Mm(24), Mm(24), Mm(24), Mm(24)]
    _format_header_row(t_geo.rows[0])
    geo_headers = [
        "№ точки",
        "С.Ш., град.",
        "С.Ш., мин.",
        "С.Ш., сек.",
        "В.Д., град.",
        "В.Д., мин.",
        "В.Д., сек.",
    ]
    for i, h in enumerate(geo_headers):
        cell = t_geo.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(h).font.bold = True

    t_geo.rows[1].cells[0].paragraphs[0].text = "{%tr for p in corner_points %}"

    d_geo = t_geo.rows[2].cells
    for i in range(7):
        _set_cell_padding(d_geo[i], top=70, bottom=70, left=60, right=60)
    d_geo[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d_geo[0].paragraphs[0].add_run("{{ p.n }}")
    d_geo[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[1].paragraphs[0].add_run("{{ p.lat_d }}")
    d_geo[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[2].paragraphs[0].add_run("{{ p.lat_m }}")
    d_geo[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[3].paragraphs[0].add_run("{{ p.lat_s }}")
    d_geo[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[4].paragraphs[0].add_run("{{ p.lon_d }}")
    d_geo[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[5].paragraphs[0].add_run("{{ p.lon_m }}")
    d_geo[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d_geo[6].paragraphs[0].add_run("{{ p.lon_s }}")

    t_geo.rows[3].cells[0].paragraphs[0].text = "{%tr endfor %}"
    _set_col_widths(t_geo, geo_widths)

    doc.add_heading("Раздел 2. Организация-недропользователь", level=1)
    t_org = doc.add_table(rows=4, cols=2)
    t_org.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_org, style="grid")
    org_widths = [Mm(70), Mm(90)]
    org_data = [
        ("Название организации", "{{ company_name }}"),
        ("Контракт/лицензия на недропользование материнского месторождения и (или) использование пространства недр", "{{ fmt.contract_info }}"),
        ("Юридический адрес", "{{ fmt.legal_address }}"),
        ("Телефон", "{{ fmt.phone }}"),
    ]
    for idx, (lbl, val) in enumerate(org_data):
        row = t_org.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        p0 = row.cells[0].paragraphs[0]
        p0.add_run(lbl).font.bold = True
        p1 = row.cells[1].paragraphs[0]
        p1.add_run(val)
    _set_col_widths(t_org, org_widths)

    doc.add_heading("Раздел 3. Способ отработки вскрышных и вмещающих пород материнского месторождения", level=1)
    t_meth = doc.add_table(rows=2, cols=2)
    t_meth.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_meth, style="grid")
    meth_widths = [Mm(80), Mm(80)]
    meth_data = [
        ("Название материнского месторождения", "{{ fmt.parent_deposit_name }}"),
        ("Способ отработки общераспространённых полезных ископаемых", "{{ fmt.mining_method }}"),
    ]
    for idx, (lbl, val) in enumerate(meth_data):
        row = t_meth.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        row.cells[0].paragraphs[0].add_run(lbl).font.bold = True
        row.cells[1].paragraphs[0].add_run(val)
    _set_col_widths(t_meth, meth_widths)

    doc.add_heading("Раздел 4. Административное положение объекта", level=1)
    t_adm = doc.add_table(rows=2, cols=3)
    t_adm.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_adm, style="grid")
    adm_widths = [Mm(55), Mm(55), Mm(50)]
    _format_header_row(t_adm.rows[0])
    adm_headers = ["Область", "Район", "Город / посёлок"]
    for i, h in enumerate(adm_headers):
        cell = t_adm.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(h).font.bold = True
    row_adm = t_adm.rows[1]
    for i, val in enumerate(["{{ fmt.oblast }}", "{{ fmt.district }}", "{{ fmt.locality }}"]):
        cell = row_adm.cells[i]
        _set_cell_padding(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(val)
    _set_col_widths(t_adm, adm_widths)

    doc.add_heading("Раздел 5. Параметры складирования вскрышных и вмещающих пород", level=1)
    t_st = doc.add_table(rows=6, cols=3)
    t_st.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_st, style="grid")
    st_widths = [Mm(85), Mm(45), Mm(30)]
    _format_header_row(t_st.rows[0])
    st_headers = ["Параметр", "Значение", "Единица"]
    for i, h in enumerate(st_headers):
        cell = t_st.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 2 else (WD_ALIGN_PARAGRAPH.RIGHT if i == 1 else WD_ALIGN_PARAGRAPH.LEFT)
        p.add_run(h).font.bold = True

    st_rows = [
        ("Длина", "{{ fmt.storage_length_m }}", "м"),
        ("Ширина", "{{ fmt.storage_width_m }}", "м"),
        ("Высота", "{{ fmt.storage_height_m }}", "м"),
        ("Площадь", "{{ fmt.storage_area_km2 }}", "км²"),
        ("Объёмы", "{{ fmt.storage_volume_thousand_m3 }}", "тыс. м³"),
    ]
    for idx, (param, val, unit) in enumerate(st_rows, start=1):
        row = t_st.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[2], top=80, bottom=80, left=100, right=100)
        row.cells[0].paragraphs[0].add_run(param)
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(val)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(unit)
    _set_col_widths(t_st, st_widths)

    doc.add_heading("Раздел 6. Методика подсчёта запасов", level=1)
    t_met = doc.add_table(rows=1, cols=1)
    t_met.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_met, style="box")
    _set_cell_padding(t_met.rows[0].cells[0], top=120, bottom=120, left=140, right=140)
    p_met = t_met.rows[0].cells[0].paragraphs[0]
    p_met.add_run("{{ fmt.reserves_methodology }}")
    _set_col_widths(t_met, [Mm(160)])

    doc.add_heading("Раздел 7. Характеристика общераспространённых полезных ископаемых", level=1)
    p_71 = doc.add_paragraph()
    r_71 = p_71.add_run("7.1. Основные характеристики:")
    r_71.font.bold = True

    p_c1 = doc.add_paragraph()
    p_c1.add_run("Минеральный состав: ").font.bold = True
    p_c1.add_run("{{ fmt.composition_mineral }}\n")
    p_c1.add_run("Петрографический и литологический состав: ").font.bold = True
    p_c1.add_run("{{ fmt.composition_petrographic }}\n")
    p_c1.add_run("Гранулометрический состав: ").font.bold = True
    p_c1.add_run("{{ fmt.composition_granulometric }}\n")
    p_c1.add_run("Физико-механические свойства: ").font.bold = True
    p_c1.add_run("{{ fmt.composition_physico_mechanical }}")

    p_ox_title = doc.add_paragraph()
    p_ox_title.paragraph_format.space_before = Pt(4)
    p_ox_title.paragraph_format.space_after = Pt(2)
    p_ox_title.add_run("Химический состав основных окислов:").font.italic = True

    t_ox = doc.add_table(rows=4, cols=2)
    t_ox.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_ox, style="grid")
    ox_widths = [Mm(100), Mm(60)]
    _format_header_row(t_ox.rows[0])
    _set_cell_background(t_ox.rows[0].cells[0], "F0F4F8")
    _set_cell_background(t_ox.rows[0].cells[1], "F0F4F8")
    _set_cell_padding(t_ox.rows[0].cells[0], top=80, bottom=80, left=100, right=100)
    _set_cell_padding(t_ox.rows[0].cells[1], top=80, bottom=80, left=100, right=100)
    p0 = t_ox.rows[0].cells[0].paragraphs[0]
    p0.add_run("Окисел").font.bold = True
    p1 = t_ox.rows[0].cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.add_run("Содержание, %").font.bold = True

    t_ox.rows[1].cells[0].paragraphs[0].text = "{%tr for ox in chemical_oxides %}"
    _set_cell_padding(t_ox.rows[2].cells[0], top=70, bottom=70, left=100, right=100)
    _set_cell_padding(t_ox.rows[2].cells[1], top=70, bottom=70, left=100, right=100)
    t_ox.rows[2].cells[0].paragraphs[0].add_run("{{ ox.name }}")
    p_ox_val = t_ox.rows[2].cells[1].paragraphs[0]
    p_ox_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ox_val.add_run("{{ ox.value_pct }}")
    t_ox.rows[3].cells[0].paragraphs[0].text = "{%tr endfor %}"
    _set_col_widths(t_ox, ox_widths)

    p_72 = doc.add_paragraph()
    p_72.paragraph_format.space_before = Pt(8)
    p_72.paragraph_format.space_after = Pt(2)
    r_72 = p_72.add_run("7.2. Дополнительные характеристики:")
    r_72.font.bold = True

    t_ex = doc.add_table(rows=4, cols=2)
    t_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_ex, style="grid")
    ex_widths = [Mm(70), Mm(90)]
    _format_header_row(t_ex.rows[0])
    _set_cell_background(t_ex.rows[0].cells[0], "F0F4F8")
    _set_cell_background(t_ex.rows[0].cells[1], "F0F4F8")
    _set_cell_padding(t_ex.rows[0].cells[0], top=80, bottom=80, left=100, right=100)
    _set_cell_padding(t_ex.rows[0].cells[1], top=80, bottom=80, left=100, right=100)
    t_ex.rows[0].cells[0].paragraphs[0].add_run("Характеристика").font.bold = True
    t_ex.rows[0].cells[1].paragraphs[0].add_run("Описание").font.bold = True

    t_ex.rows[1].cells[0].paragraphs[0].text = "{%tr for ch in extra_characteristics %}"
    _set_cell_padding(t_ex.rows[2].cells[0], top=70, bottom=70, left=100, right=100)
    _set_cell_padding(t_ex.rows[2].cells[1], top=70, bottom=70, left=100, right=100)
    t_ex.rows[2].cells[0].paragraphs[0].add_run("{{ ch.name }}")
    t_ex.rows[2].cells[1].paragraphs[0].add_run("{{ ch.description }}")
    t_ex.rows[3].cells[0].paragraphs[0].text = "{%tr endfor %}"
    _set_col_widths(t_ex, ex_widths)

    doc.add_heading("Раздел 8. Область применения", level=1)
    t_app = doc.add_table(rows=4, cols=3)
    t_app.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_app, style="grid")
    app_widths = [Mm(50), Mm(60), Mm(50)]
    _format_header_row(t_app.rows[0])
    app_headers = ["Полезное ископаемое", "Область применения", "Соответствие ГОСТу или СТ РК"]
    for i, h in enumerate(app_headers):
        cell = t_app.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=100, right=100)
        cell.paragraphs[0].add_run(h).font.bold = True

    t_app.rows[1].cells[0].paragraphs[0].text = "{%tr for app in applications %}"
    _set_cell_padding(t_app.rows[2].cells[0], top=70, bottom=70, left=100, right=100)
    _set_cell_padding(t_app.rows[2].cells[1], top=70, bottom=70, left=100, right=100)
    _set_cell_padding(t_app.rows[2].cells[2], top=70, bottom=70, left=100, right=100)
    t_app.rows[2].cells[0].paragraphs[0].add_run("{{ app.mineral }}")
    t_app.rows[2].cells[1].paragraphs[0].add_run("{{ app.usage }}")
    t_app.rows[2].cells[2].paragraphs[0].add_run("{{ app.standard }}")
    t_app.rows[3].cells[0].paragraphs[0].text = "{%tr endfor %}"
    _set_col_widths(t_app, app_widths)

    doc.add_heading("Раздел 9. Запасы", level=1)
    t_res = doc.add_table(rows=6, cols=2)
    t_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_res, style="grid")
    res_widths = [Mm(100), Mm(60)]
    res_data = [
        ("Тип полезного ископаемого", "{{ fmt.rock_origin_text }}"),
        ("Полезное ископаемое и единица измерения", "{{ mineral_type }}, {{ fmt.unit }}"),
        ("Запасы по предыдущему паспорту", "№ {{ fmt.previous_passport_number }}: {{ fmt.previous_reserves }} {{ fmt.unit }}"),
        ("Прирост запасов в связи со складированием", "{{ fmt.reserves_increment }} {{ fmt.unit }}"),
        ("Реализованные запасы", "{{ fmt.reserves_realized }} {{ fmt.unit }}"),
        ("Запасы на дату заполнения паспорта", "{{ fmt.current_reserves }} {{ fmt.unit }} (на дату {{ fmt.passport_date }})"),
    ]
    for idx, (lbl, val) in enumerate(res_data):
        row = t_res.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        row.cells[0].paragraphs[0].add_run(lbl).font.bold = True
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.add_run(val)
    _set_col_widths(t_res, res_widths)

    doc.add_heading("Приложения", level=1)

    p_app1 = doc.add_paragraph()
    p_app1.paragraph_format.space_before = Pt(8)
    p_app1.paragraph_format.space_after = Pt(4)
    r_a1 = p_app1.add_run("Приложение 1. Пояснительная записка")
    r_a1.font.bold = True
    r_a1.font.size = Pt(12)

    doc.add_paragraph("{%p for p in explanatory_note %}")
    p_note_body = doc.add_paragraph()
    p_note_body.paragraph_format.line_spacing = 1.15
    p_note_body.paragraph_format.space_after = Pt(4)
    p_note_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_note_body.add_run("{{ p }}")
    doc.add_paragraph("{%p endfor %}")

    p_app2 = doc.add_paragraph()
    p_app2.paragraph_format.space_before = Pt(14)
    p_app2.paragraph_format.space_after = Pt(4)
    r_a2 = p_app2.add_run("Приложение 2. Графические материалы")
    r_a2.font.bold = True
    r_a2.font.size = Pt(12)

    t_graph = doc.add_table(rows=5, cols=3)
    t_graph.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(t_graph, style="grid")
    graph_widths = [Mm(15), Mm(95), Mm(50)]
    _format_header_row(t_graph.rows[0])
    graph_headers = ["№", "Наименование графического приложения", "Статус"]
    for i, h in enumerate(graph_headers):
        cell = t_graph.rows[0].cells[i]
        _set_cell_background(cell, "F0F4F8")
        _set_cell_padding(cell, top=80, bottom=80, left=100, right=100)
        cell.paragraphs[0].add_run(h).font.bold = True

    graph_rows = [
        ("1", "Ситуационный план объекта", "прилагается отдельно"),
        ("2", "Геологическая карта объекта", "прилагается отдельно"),
        ("3", "План подсчёта запасов полезного ископаемого", "прилагается отдельно"),
        ("4", "Геологические разрезы объекта", "прилагается отдельно"),
    ]
    for idx, (num, name, stat) in enumerate(graph_rows, start=1):
        row = t_graph.rows[idx]
        _set_cell_padding(row.cells[0], top=80, bottom=80, left=80, right=80)
        _set_cell_padding(row.cells[1], top=80, bottom=80, left=100, right=100)
        _set_cell_padding(row.cells[2], top=80, bottom=80, left=100, right=100)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.add_run(num)
        row.cells[1].paragraphs[0].add_run(name)
        p2 = row.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_stat = p2.add_run(stat)
        r_stat.font.italic = True
        r_stat.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    _set_col_widths(t_graph, graph_widths)

    p_footer = doc.add_paragraph()
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_footer.paragraph_format.space_before = Pt(20)
    r_f1 = p_footer.add_run("Документ сформирован автоматически сервисом GeoSaaS API\n")
    r_f1.font.italic = True
    r_f1.font.size = Pt(9)
    r_f1.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    r_f2 = p_footer.add_run("Дата формирования: {{ generated_date }}")
    r_f2.font.italic = True
    r_f2.font.size = Pt(9)
    r_f2.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


if __name__ == "__main__":
    target = Path(__file__).resolve().parent / "templates" / "passport_template.docx"
    generate_template(target)
